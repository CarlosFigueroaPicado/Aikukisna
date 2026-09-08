from pathlib import Path
from datetime import date
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "aikukisna"
OUTPUT = DOCS / "Documentacion_Tecnica_Aikukisna.docx"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="1F2937", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        set_cell_text(table.rows[0].cells[index], value, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], "0B6E99")
        set_cell_border(table.rows[0].cells[index])
    for row_index, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if row_index % 2 == 1:
                shade(cells[index], "F2F7FA")
            set_cell_border(cells[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, text):
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(text.rstrip())


def add_markdown_file(doc, path, section_break=False):
    if section_break:
        doc.add_page_break()
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines = []
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            flush_table()
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r"[-: ]+", cell or "-") for cell in cells):
                index += 1
                continue
            table_rows.append(cells)
            index += 1
            continue
        flush_table()
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            paragraph.add_run(re.sub(r"`([^`]+)`", r"\1", heading.group(2)))
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(re.sub(r"`([^`]+)`", r"\1", bullet.group(1)))
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.add_run(re.sub(r"`([^`]+)`", r"\1", numbered.group(1)))
            index += 1
            continue
        paragraph = doc.add_paragraph(style="Body Text")
        paragraph.add_run(re.sub(r"`([^`]+)`", r"\1", line))
        index += 1
    flush_table()


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for level, size, color in [(1, 18, "0B6E99"), (2, 14, "0B6E99"), (3, 11.5, "168AAD"), (4, 10.5, "374151")]:
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if level < 3 else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.2)
    code.font.color.rgb = RGBColor(31, 41, 55)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AIKUKISNA  |  DOCUMENTACION TECNICA")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(107, 114, 128)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("YAWANSA TECH  |  Uso académico y técnico")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(107, 114, 128)


def main():
    doc = Document()
    configure_document(doc)
    add_header_footer(doc.sections[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Documentación Técnica Aikukisna")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(11, 110, 153)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Arquitectura, base de datos, interfaz, seguridad y despliegue")
    run.font.name = "Aptos"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Aplicación Android para el aprendizaje de idiomas dirigida a estudiantes de zonas rurales.").italic = True
    doc.add_paragraph()
    date_paragraph = doc.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.add_run(date.today().strftime("%d de %B de %Y"))

    doc.add_page_break()
    doc.add_heading("Contenido", level=1)
    toc_items = [
        "1. README y documentación técnica",
        "2. Diagramación de base de datos",
        "3. Interfaz y desarrollo",
        "4. Control de versiones",
        "5. Seguridad y buenas prácticas",
        "6. Ejecución de la solución",
        "Anexo. Scripts y evidencias de endpoints",
    ]
    for item in toc_items:
        doc.add_paragraph(item, style="List Number")
    doc.add_paragraph("El documento refleja la implementación disponible en la rama de trabajo y distingue las configuraciones externas de Supabase, Google, Gemini y ElevenLabs.")

    files = [
        DOCS / "01-readme-tecnico.md",
        DOCS / "02-base-de-datos.md",
        DOCS / "Documentacion_interfaz_y_desarrollo.md",
        DOCS / "04-control-de-versiones.md",
        DOCS / "05-seguridad-y-buenas-practicas.md",
        DOCS / "06-ejecucion.md",
    ]
    for path in files:
        add_markdown_file(doc, path, section_break=True)

    doc.add_page_break()
    doc.add_heading("Anexo Scripts y evidencias de endpoints", level=1)
    doc.add_paragraph("El proyecto incluye scripts reproducibles para validar variables, consultar Supabase, generar evidencia y construir el APK.")
    add_code(doc, ".\\scripts\\verify-environment.ps1\n.\\scripts\\verify-supabase.ps1\n.\\scripts\\demo-endpoints.ps1\n.\\scripts\\verify-build.ps1\n.\\scripts\\deploy-release.ps1")
    doc.add_heading("Resultado de evidencia", level=2)
    evidence = DOCS / "evidencias" / "resumen.json"
    if evidence.exists():
        add_code(doc, evidence.read_text(encoding="utf-8"))
    doc.add_heading("Archivos de configuración", level=2)
    add_code(doc, "local.properties.example\ngradle/libs.versions.toml\napp/build.gradle.kts\ngradlew\ngradlew.bat")

    for section in doc.sections:
        add_header_footer(section)
    doc.core_properties.title = "Documentación Técnica Aikukisna"
    doc.core_properties.subject = "Arquitectura, desarrollo y despliegue de Aikukisna"
    doc.core_properties.author = "YAWANSA TECH"
    doc.core_properties.keywords = "Aikukisna, Android, Supabase, documentación técnica"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
